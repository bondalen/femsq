#!/usr/bin/env python3
"""Excel-макет Rslt для песочницы test_sudz_sm: настоящие вертикальные merge."""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HERE = Path(__file__).resolve().parent
OUT_NAME = "ags_Yr_DbtChangesRslt_sm_sandbox.xlsx"

QUARTER_FILLS = ["D9D9D9", "D6D4CB", "FFFFCC", "E6E0EC", "B7DEE8"]
FILL_BASE = "FDEADA"
FILL_INV = "D99694"
FILL_OVERD = "FFFF00"
FILL_CURATOR = "D7E4BD"
FILL_NEW = "F2DCDB"
FILL_TECH = "C0C0C0"
FILL_SUM_TTL = "D9D9D9"
FILL_SUM_POG = "FAC090"
FILL_README = "FFF8E7"
FONT_SUM = "C0504D"
MONEY_FMT = '#,##0.00'
ROMAN = ["I", "II", "III", "IV"]

THIN = Border(
    left=Side(style="thin", color="808080"),
    right=Side(style="thin", color="808080"),
    top=Side(style="thin", color="808080"),
    bottom=Side(style="thin", color="808080"),
)
FONT_H = Font(name="Calibri", size=8)
FONT_T = Font(name="Calibri", size=9, bold=True)
FONT_D = Font(name="Calibri", size=9)
FONT_S = Font(name="Calibri", size=8, bold=True, color=FONT_SUM)
FONT_TITLE = Font(name="Calibri", size=14, bold=True)
ALIGN_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_L = Alignment(horizontal="left", vertical="center", wrap_text=True)


@dataclass(frozen=True)
class UplMeta:
    key: int
    as_of: date
    name: str


@dataclass(frozen=True)
class Val:
    upl: int
    slot: int
    dv_key: int
    ttl: Decimal
    overd: Decimal
    id_num: int
    inv_num: str
    cn_num: str


@dataclass(frozen=True)
class Cmm:
    dv_key: int
    gr: int
    text: str
    deep: bool
    slot: int


@dataclass
class Col:
    human: str
    tech: str
    kind: str  # key, text, money, pog
    fill: str | None
    width: int


def _fill(hex6: str | None) -> PatternFill | None:
    if not hex6:
        return None
    return PatternFill("solid", fgColor=hex6)


def _q_label(d: date) -> str:
    q = (d.month - 1) // 3
    return f"{d.year}. {ROMAN[q]}-й квартал"


def _as_date(v: Any) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return date.fromisoformat(str(v)[:10])


def _dec(v: Any) -> Decimal:
    return Decimal("0") if v is None else Decimal(str(v))


def _money(v: Decimal | None) -> float | None:
    if v is None:
        return None
    return float(v)


def _pogasheno(base: Decimal, curr: Decimal) -> Decimal | None:
    delta = base - curr
    if delta <= 0:
        return None
    return delta


def load_sandbox(cur) -> tuple[dict[int, UplMeta], dict[int, dict[int, list[Val]]], dict[tuple[int, int], list[Cmm]]]:
    """Читает upl / Value / cmm из test_sudz_sm."""
    cur.execute(
        "SELECT upl_key, uplStatusOnDate, upl_name FROM test_sudz_sm.upl ORDER BY upl_key"
    )
    upls = {int(r[0]): UplMeta(int(r[0]), _as_date(r[1]), r[2]) for r in cur.fetchall()}

    cur.execute(
        """
        SELECT
            idd.iddDbt, dv.dvUpl, dv.dvInvDbt, dv.dvKey,
            dv.dvTtl, dv.dvOverd, id.idNum, inv.iNum, cnn.cnnNum
        FROM test_sudz_sm.DbtValue AS dv
        JOIN test_sudz_sm.invDbt AS id ON id.idKey = dv.dvInvDbt
        JOIN test_sudz_sm.inv ON inv.iKey = id.idInv
        JOIN test_sudz_sm.invDbtDbt AS idd ON idd.iddInvDbt = dv.dvInvDbt
        LEFT JOIN test_sudz_sm.invDbtVar AS v ON v.idvvKey = dv.dvInvDbtVar
        LEFT JOIN test_sudz_sm.cnNum AS cnn ON cnn.cnnKey = v.idvvCnNum
        ORDER BY idd.iddDbt, dv.dvUpl, dv.dvInvDbt
        """
    )
    facts: dict[int, dict[int, list[Val]]] = defaultdict(lambda: defaultdict(list))
    for dbt, upl, slot, dv_key, ttl, overd, id_num, inv_num, cn_num in cur.fetchall():
        facts[int(dbt)][int(upl)].append(
            Val(
                int(upl),
                int(slot),
                int(dv_key),
                _dec(ttl),
                _dec(overd),
                int(id_num),
                inv_num or "",
                cn_num or "",
            )
        )

    cur.execute(
        """
        SELECT c.cmmDv, c.cmmGr, c.cmmText, dv.dvInvDbt, idd.iddDbt,
               CASE WHEN gr.cnigKey IS NULL THEN 0 ELSE 1 END
        FROM test_sudz_sm.cmm AS c
        JOIN test_sudz_sm.DbtValue AS dv ON dv.dvKey = c.cmmDv
        JOIN test_sudz_sm.invDbtDbt AS idd ON idd.iddInvDbt = dv.dvInvDbt
        LEFT JOIN test_sudz_sm.cnInvGr AS gr
            ON gr.cnigDv = c.cmmDv AND gr.cnigCmmGr = c.cmmGr
        WHERE c.cmmKind = N'mery'
        ORDER BY idd.iddDbt, c.cmmGr, dv.dvInvDbt
        """
    )
    cmms: dict[tuple[int, int], list[Cmm]] = defaultdict(list)
    for dv_key, gr, text, slot, dbt, deep in cur.fetchall():
        cmms[(int(dbt), int(gr))].append(
            Cmm(int(dv_key), int(gr), text or "", bool(deep), int(slot))
        )
    return upls, facts, cmms


def _paint(cell, *, font: Font, fill: str | None, align: Alignment, money: bool = False) -> None:
    cell.font = font
    cell.alignment = align
    cell.border = THIN
    if fill:
        cell.fill = _fill(fill)
    if money:
        cell.number_format = MONEY_FMT


def _write_header(ws: Worksheet, cols: list[Col]) -> None:
    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 72
    ws.row_dimensions[3].height = 22
    for i, col in enumerate(cols, start=1):
        sum_cell = ws.cell(1, i)
        if col.kind in ("money", "pog"):
            letter = get_column_letter(i)
            sum_cell.value = f"=SUBTOTAL(9,{letter}4:{letter}5)"
            _paint(sum_cell, font=FONT_S, fill=FILL_SUM_POG if col.kind == "pog" else FILL_SUM_TTL, align=ALIGN_C, money=True)
        else:
            _paint(sum_cell, font=FONT_H, fill=None, align=ALIGN_C)
        h = ws.cell(2, i, col.human)
        _paint(h, font=FONT_H, fill=col.fill, align=ALIGN_C)
        t = ws.cell(3, i, col.tech)
        _paint(t, font=FONT_T, fill=FILL_TECH, align=ALIGN_C)
        ws.column_dimensions[get_column_letter(i)].width = col.width


def _merge_col(ws: Worksheet, col: int, r1: int, r2: int, value: Any, *, fill: str | None, money: bool, comment: str | None = None) -> None:
    align = ALIGN_C if money or not isinstance(value, str) else ALIGN_L
    top = ws.cell(r1, col, value if not isinstance(value, Decimal) else _money(value))
    _paint(top, font=FONT_D, fill=fill, align=align, money=money)
    if comment:
        top.comment = Comment(comment, "test_sudz_sm")
    for r in range(r1 + 1, r2 + 1):
        cell = ws.cell(r, col)
        _paint(cell, font=FONT_D, fill=fill, align=align, money=money)
    if r2 > r1:
        ws.merge_cells(start_row=r1, start_column=col, end_row=r2, end_column=col)


def _put_items(
    ws: Worksheet,
    col: int,
    row0: int,
    n_rows: int,
    items: list[Any],
    *,
    fill: str | None,
    money: bool,
    merge_note: str | None = None,
) -> None:
    """Одна величина → merge на n_rows; несколько → по строке без merge."""
    if len(items) <= 1 and n_rows > 1:
        value = items[0] if items else None
        note = merge_note or "слито: зерно колонки = 1 Value, строк полосы больше"
        _merge_col(ws, col, row0, row0 + n_rows - 1, value, fill=fill, money=money, comment=note)
        return
    for i in range(n_rows):
        value = items[i] if i < len(items) else None
        if isinstance(value, Decimal):
            value = _money(value)
        cell = ws.cell(row0 + i, col, value)
        is_text = isinstance(items[i] if i < len(items) else None, str)
        _paint(cell, font=FONT_D, fill=fill, align=ALIGN_L if is_text else ALIGN_C, money=money)


def _slice_cols(upl: UplMeta, index: int, first: bool) -> list[Col]:
    q = _q_label(upl.as_of)
    p = upl.as_of.isoformat()
    band = QUARTER_FILLS[index % 5]
    cols = [
        Col("Реквизиты документа основания (СФ)", f"{p}_invNumEnum", "text", FILL_INV, 16),
        Col(f"{q}. № задолженности в СФ", f"{p}_idNum", "text", band, 12),
        Col(f"{q}. Договор", f"{p}_cnNumEnum", "text", band, 12),
        Col(f"{q}. Общая задолженность", f"{p}_Ttl", "money", band, 14),
        Col(f"{q}. Просроченная задолженность", f"{p}_Overd", "money", FILL_OVERD, 14),
    ]
    if not first:
        cols.append(Col(f"{q}. Погашенная задолженность", f"{p}_погашено", "pog", band, 14))
    return cols


def write_band(
    wb: Workbook,
    sheet_name: str,
    caption: str,
    dbt: int,
    fact_upls: list[int],
    comment_specs: list[tuple[int, str, str, str | None]],
    upls: dict[int, UplMeta],
    facts: dict[int, dict[int, list[Val]]],
    cmms: dict[tuple[int, int], list[Cmm]],
) -> None:
    """
    Пишет лист Rslt-полосы.

    comment_specs: (cmmGr, human, tech, fill) — fill None → CURATOR, иначе NEW.
    «погашено» всегда зерно Dbt (одна слитая ячейка), чтобы не размножать 18000×2.
    """
    ws = wb.create_sheet(sheet_name)
    by_upl = facts.get(dbt, {})
    n_rows = 1
    for u in fact_upls:
        n_rows = max(n_rows, len(by_upl.get(u, [])))
    for gr, *_rest in comment_specs:
        n_rows = max(n_rows, len(cmms.get((dbt, gr), [])))

    cols: list[Col] = [
        Col("", "dbtKey", "key", None, 10),
        Col("Счет Главной книги", "account_num", "text", FILL_BASE, 14),
    ]
    for i, u in enumerate(fact_upls):
        cols.extend(_slice_cols(upls[u], i, first=(i == 0)))
    for gr, human, tech, fill in comment_specs:
        cols.append(Col(human, tech, "text", fill or FILL_CURATOR, 36))
        cols.append(Col("Группа («углубленно»)", f"gr_{gr}", "text", fill or FILL_CURATOR, 16))

    _write_header(ws, cols)
    # SUBTOTAL диапазон под фактическое число строк
    last = 3 + n_rows
    for i, col in enumerate(cols, start=1):
        if col.kind in ("money", "pog"):
            letter = get_column_letter(i)
            ws.cell(1, i).value = f"=SUBTOTAL(9,{letter}4:{letter}{last})"

    row0 = 4
    for r in range(row0, row0 + n_rows):
        ws.row_dimensions[r].height = 36

    col_i = 1
    _put_items(ws, col_i, row0, n_rows, [dbt], fill=None, money=False, merge_note="канон Dbt на всю полосу")
    col_i += 1
    _put_items(ws, col_i, row0, n_rows, [24], fill=FILL_BASE, money=False)
    col_i += 1

    base_overd = sum((v.overd for v in by_upl.get(fact_upls[0], [])), Decimal("0")) if fact_upls else Decimal("0")

    for si, u in enumerate(fact_upls):
        vals = by_upl.get(u, [])
        first = si == 0
        qfill = QUARTER_FILLS[si % 5]
        _put_items(ws, col_i, row0, n_rows, [v.inv_num for v in vals], fill=FILL_INV, money=False)
        col_i += 1
        _put_items(ws, col_i, row0, n_rows, [v.id_num for v in vals], fill=qfill, money=False)
        col_i += 1
        _put_items(ws, col_i, row0, n_rows, [v.cn_num for v in vals], fill=qfill, money=False)
        col_i += 1
        _put_items(ws, col_i, row0, n_rows, [v.ttl for v in vals], fill=qfill, money=True)
        col_i += 1
        _put_items(ws, col_i, row0, n_rows, [v.overd for v in vals], fill=FILL_OVERD, money=True)
        col_i += 1
        if not first:
            curr_overd = sum((v.overd for v in vals), Decimal("0"))
            pog = _pogasheno(base_overd, curr_overd)
            _put_items(
                ws,
                col_i,
                row0,
                n_rows,
                [pog],
                fill=qfill,
                money=True,
                merge_note="погашено на зерне Dbt (∑ Overd), не на каждой доле",
            )
            col_i += 1

    for gr, _human, _tech, fill in comment_specs:
        bunch = cmms.get((dbt, gr), [])
        cfill = fill or FILL_CURATOR
        _put_items(ws, col_i, row0, n_rows, [c.text for c in bunch], fill=cfill, money=False)
        col_i += 1
        _put_items(
            ws,
            col_i,
            row0,
            n_rows,
            [("углубленно" if c.deep else "") for c in bunch],
            fill=cfill,
            money=False,
        )
        col_i += 1

    ws.freeze_panes = "C4"
    ws.auto_filter.ref = f"B2:{get_column_letter(len(cols))}{last}"
    ws.sheet_view.showGridLines = True
    ws.oddHeader.left.text = caption
    # подпись над таблицей в комментарии A2, чтобы не ломать шапку Access
    ws.cell(2, 1).comment = Comment(caption, "test_sudz_sm")


def write_readme(wb: Workbook) -> None:
    ws = wb.create_sheet("readme", 0)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 110
    lines = [
        ("Макет", "Песочница split/merge test_sudz_sm — не выгрузка Access и не прод-экспорт FEMSQ."),
        ("Зачем", "Показать вертикальный merge, когда зерно факта и зерно комментария различаются."),
        ("Правило", "Факт квартала = число DbtValue на этой upl. Комментарии = число Value среза своей группы."),
        ("Погашено", "Одна слитая ячейка на канон (∑ Overd). Не писать 18000 на каждую долю — это ложное «погашено»."),
        ("Лист A", "Крайний срез 202 = 2 строки факта; старые cmm с 201 = одна ячейка на обе строки."),
        ("Лист B", "Крайний срез 201 = одна ячейка факта; «старые» cmm с 202 = две ячейки (взгляд назад)."),
        ("Лист C", "После merge-back: крайний 210 = 36000 слито; старые cmm с 209 = две доли."),
        ("Лист 112", "Слияние канонов 113→112: две Value / два комментария, merge не нужен."),
        ("Лист 111_timeline", "Все срезы 201–210 на одной полосе из двух строк."),
        ("Жёлтый", "Просрочка. Зелёный — мероприятия. Розовый — *_new. Серый SUBTOTAL — суммы колонки."),
        ("Комментарий ячейки", "У слитых ячеек жёлтый уголок: почему merge."),
        ("Прод", "SudzRsltExcelExporter пока пишет 1 строку на Dbt. Этот файл — целевое поведение полосы."),
    ]
    ws.cell(1, 1, "test_sudz_sm · макет Rslt").font = FONT_TITLE
    ws.merge_cells("A1:B1")
    ws.row_dimensions[1].height = 24
    for i, (k, v) in enumerate(lines, start=3):
        a = ws.cell(i, 1, k)
        b = ws.cell(i, 2, v)
        a.font = FONT_T
        b.font = FONT_D
        a.fill = _fill(FILL_README)
        b.fill = _fill(FILL_README)
        a.alignment = ALIGN_L
        b.alignment = ALIGN_L
        ws.row_dimensions[i].height = 20


def write_sandbox_xlsx(cur, out_dir: Path | None = None) -> Path:
    """Собирает книгу макета и пишет её рядом со скриптами песочницы."""
    upls, facts, cmms = load_sandbox(cur)
    wb = Workbook()
    # default sheet replaced by readme
    default = wb.active
    wb.remove(default)
    write_readme(wb)

    write_band(
        wb,
        "A_111_curr202",
        "A: QI после split. Крайний 202 = 2 факта; старые мероприятия с 201 слиты на 2 строки.",
        111,
        [201, 202],
        [
            (201, "Мероприятия (старые · свод 201)", "mery", None),
            (202, "Мероприятия, новые, по состоянию на март 2025", "mery_new", FILL_NEW),
        ],
        upls,
        facts,
        cmms,
    )
    write_band(
        wb,
        "B_111_curr201",
        "B: крайний 201 = 1 факт (слит на 2 строки); комментарии группы 202 = две ячейки.",
        111,
        [201],
        [
            (202, "Мероприятия (группа 202, как «старые» при взгляде с 201)", "mery", None),
        ],
        upls,
        facts,
        cmms,
    )
    write_band(
        wb,
        "C_111_curr210",
        "C: merge-back. Крайний 210 = снова 36000 слито; старые с 209 = две доли без merge.",
        111,
        [209, 210],
        [
            (209, "Мероприятия (старые · свод 209, ещё split)", "mery", None),
            (210, "Мероприятия, новые, по состоянию на март 2027", "mery_new", FILL_NEW),
        ],
        upls,
        facts,
        cmms,
    )
    write_band(
        wb,
        "112_canon_merge",
        "112 после UPDATE моста 113→112: две Value на срезе 201 и два комментария — merge нет.",
        112,
        [201],
        [
            (201, "Мероприятия (свод 201)", "mery", None),
        ],
        upls,
        facts,
        cmms,
    )
    write_band(
        wb,
        "111_timeline",
        "111 по всем срезам 201–210. Серые/жёлтые блоки с одной Value слиты; 202–209 — по две строки.",
        111,
        list(range(201, 211)),
        [
            (201, "mery 201 (целый)", "mery_201", None),
            (202, "mery 202 (две доли)", "mery_202", FILL_NEW),
            (209, "mery 209 (ещё две доли)", "mery_209", None),
            (210, "mery 210 (снова целый)", "mery_210", FILL_NEW),
        ],
        upls,
        facts,
        cmms,
    )

    out = (out_dir or HERE) / OUT_NAME
    wb.save(out)
    return out
