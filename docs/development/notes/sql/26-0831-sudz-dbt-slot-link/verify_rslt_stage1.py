#!/usr/bin/env python3
"""
Stage1 / E1′ gate: FEMSQ Rslt (база–QI–QII) vs Access ags_Yr_DbtChangesRslt_*.xlsx.

Игнорирует колонки QIII/QIV в эталоне (отдельного Excel QIII нет).

Usage:
  python3 verify_rslt_stage1.py --ref PATH.xlsx --gen PATH.xlsx [--sample 80] [--json OUT.json]
  # exit 0 = PASS, 1 = FAIL

Колонки Access (с ciaKey): base K/L; QI Y/Z/AF; QII AN/AO/AU; SF QI=T(20), QII=AI(35).
Колонки FEMSQ (без ciaKey): те же Ttl/Overd индексы для base/QI/QII Ttl;
  SF QI=17, QII=32; погашено QI=31, QII=46.
"""
from __future__ import annotations

import argparse
import json
import random
import sys
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("NEED openpyxl (e.g. /tmp/rslt-venv/bin/python)", file=sys.stderr)
    sys.exit(2)

def D(v):
    if v is None or v == "":
        return None
    try:
        return Decimal(str(v)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except Exception:
        return None


def money0(v):
    return D(v) if D(v) is not None else Decimal("0")


def norm_sf(v):
    if v is None or v == "":
        return None
    return str(v).strip()


def sf_ok(ref_sf, gen_sf, left_inv):
    rs, gs, li = norm_sf(ref_sf), norm_sf(gen_sf), norm_sf(left_inv)
    if rs == gs:
        return True
    # Access: пусто при неизменном СФ; FEMSQ пишет invNumEnum
    if rs is None and gs == li:
        return True
    return False


# 1-based column maps for base–QI–QII only (QIII/QIV ignored)
REF = {
    "acc": 2,
    "inv": 3,
    "cn": 5,
    "ttl": 11,
    "o0": 12,
    "sf_i": 20,
    "ttl_i": 25,
    "o_i": 26,
    "pog_i": 32,
    "sf_ii": 35,
    "ttl_ii": 40,
    "o_ii": 41,
    "pog_ii": 47,
}
GEN = {
    "acc": 2,
    "inv": 3,
    "cn": 5,
    "ttl": 11,
    "o0": 12,
    "sf_i": 17,
    "ttl_i": 25,
    "o_i": 26,
    "pog_i": 31,
    "sf_ii": 32,
    "ttl_ii": 40,
    "o_ii": 41,
    "pog_ii": 46,
}

SUM_KEYS = ("ttl", "o0", "ttl_i", "o_i", "pog_i", "ttl_ii", "o_ii", "pog_ii")


def load_rows(path: str, cols: dict) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    max_c = max(cols.values())
    out = []
    for r, row in enumerate(ws.iter_rows(min_row=4, max_col=max_c, values_only=True), start=4):
        if not row:
            continue
        rec = {"row": r}
        for name, idx in cols.items():
            rec[name] = row[idx - 1] if len(row) >= idx else None
        out.append(rec)
    wb.close()
    return out


def sum_cols(rows: list[dict], keys=SUM_KEYS):
    totals = {k: Decimal("0") for k in keys}
    nonempty = {k: 0 for k in keys}
    for rec in rows:
        for k in keys:
            if rec.get(k) not in (None, ""):
                nonempty[k] += 1
                totals[k] += money0(rec[k])
    return totals, nonempty


def biz_key(rec) -> tuple:
    return (
        str(rec.get("acc") or "").strip(),
        str(rec.get("inv") or "").strip(),
        D(rec.get("ttl")),
        str(rec.get("cn") or "").strip(),
    )


def index_by_key(rows):
    idx = defaultdict(list)
    for rec in rows:
        idx[biz_key(rec)].append(rec)
    return idx


def compare_pair(ref, gen, label: str) -> list[str]:
    errs = []
    for fld, name in (
        ("ttl", "base_Ttl"),
        ("o0", "base_Overd"),
        ("ttl_i", "q1_Ttl"),
        ("o_i", "q1_Overd"),
        ("pog_i", "q1_pog"),
        ("ttl_ii", "q2_Ttl"),
        ("o_ii", "q2_Overd"),
        ("pog_ii", "q2_pog"),
    ):
        rv, gv = D(ref.get(fld)), D(gen.get(fld))
        # both empty OK
        if ref.get(fld) in (None, "") and gen.get(fld) in (None, ""):
            continue
        if rv != gv:
            # Access empty overd/ttl vs FEMSQ absent period: treat empty==empty already;
            # empty vs 0: Access may omit 0? compare carefully
            if rv is None and gv == Decimal("0"):
                continue
            if gv is None and rv == Decimal("0"):
                continue
            errs.append(f"{label}.{name}: ref={rv} gen={gv}")
    if not sf_ok(ref.get("sf_i"), gen.get("sf_i"), ref.get("inv")):
        errs.append(f"{label}.sf_i: ref={ref.get('sf_i')!r} gen={gen.get('sf_i')!r}")
    if not sf_ok(ref.get("sf_ii"), gen.get("sf_ii"), ref.get("inv")):
        errs.append(f"{label}.sf_ii: ref={ref.get('sf_ii')!r} gen={gen.get('sf_ii')!r}")
    return errs


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--ref", required=True, help="Access Rslt xlsx")
    ap.add_argument("--gen", required=True, help="FEMSQ Rslt xlsx (asOfUpl=803)")
    ap.add_argument("--sample", type=int, default=80, help="stratified sample size")
    ap.add_argument("--seed", type=int, default=31, help="RNG seed")
    ap.add_argument("--json", default="", help="write report JSON")
    args = ap.parse_args()

    report = {"pass": True, "checks": {}, "errors": []}
    ref_rows = load_rows(args.ref, REF)
    gen_rows = load_rows(args.gen, GEN)
    report["rows"] = {"ref": len(ref_rows), "gen": len(gen_rows)}

    rt, rn = sum_cols(ref_rows)
    gt, gn = sum_cols(gen_rows)
    sum_ok = True
    sum_detail = {}
    for k in SUM_KEYS:
        ok = rt[k] == gt[k]
        sum_ok = sum_ok and ok
        sum_detail[k] = {
            "ref": str(rt[k]),
            "gen": str(gt[k]),
            "ref_n": rn[k],
            "gen_n": gn[k],
            "ok": ok,
        }
        if not ok:
            report["errors"].append(f"SUM {k}: ref={rt[k]} gen={gt[k]}")
    report["checks"]["row1_sums_base_qi_qii"] = {"ok": sum_ok, "detail": sum_detail}
    print("=== row1 SUBTOTAL base–QI–QII (QIII/QIV ignored) ===")
    for k in SUM_KEYS:
        d = sum_detail[k]
        print(f"  {k}: ref={d['ref']} (n={d['ref_n']}) gen={d['gen']} (n={d['gen_n']}) {'OK' if d['ok'] else 'DIFF'}")

    # Control keys: 7947 + L001
    controls = [
        ("7947", Decimal("70525000.01"), "1-306", "stable7947"),
        ("А19-16343/2021", Decimal("9527.42"), "32-425", "L001"),
    ]
    gen_idx = index_by_key(gen_rows)
    ctrl_errs = []
    print("=== controls ===")
    for inv, ttl, cn_part, tag in controls:
        cands = [
            r
            for r in ref_rows
            if str(r.get("inv")) == inv
            and D(r.get("ttl")) == ttl
            and cn_part in str(r.get("cn") or "")
        ]
        if not cands:
            ctrl_errs.append(f"{tag}: not in REF")
            continue
        ref = cands[0]
        g_cands = [
            g
            for key, lst in gen_idx.items()
            for g in lst
            if key[1] == inv and key[2] == ttl and cn_part in key[3]
        ]
        if not g_cands:
            ctrl_errs.append(f"{tag}: not in GEN")
            print(f"  {tag}: MISSING in GEN")
            continue
        errs = compare_pair(ref, g_cands[0], tag)
        if errs:
            ctrl_errs.extend(errs)
            print(f"  {tag}: DIFF", *errs)
        else:
            print(f"  {tag}: OK")
    report["checks"]["controls"] = {"ok": not ctrl_errs, "errors": ctrl_errs}
    report["errors"].extend(ctrl_errs)

    # All Access rows with погашено QI or QII
    print("=== погашено rows (all Access nonempty) ===")
    pog_errs = []
    pog_n = 0
    for ref in ref_rows:
        if ref.get("pog_i") in (None, "") and ref.get("pog_ii") in (None, ""):
            continue
        pog_n += 1
        key = biz_key(ref)
        # match gen loosely: acc+inv+ttl (cn may differ slightly)
        g_cands = [
            g
            for g in gen_rows
            if str(g.get("acc") or "").strip() == key[0]
            and str(g.get("inv") or "").strip() == key[1]
            and D(g.get("ttl")) == key[2]
        ]
        if not g_cands:
            msg = f"pog row{ref['row']} {key[1]}/{key[2]}: MISSING in GEN"
            pog_errs.append(msg)
            print(" ", msg)
            continue
        errs = compare_pair(ref, g_cands[0], f"pog_r{ref['row']}")
        if errs:
            pog_errs.extend(errs)
            print(f"  row{ref['row']} {key[1]}: DIFF", *errs)
        else:
            print(f"  row{ref['row']} {key[1]}: OK")
    report["checks"]["pogasheno_rows"] = {"ok": not pog_errs, "n": pog_n, "errors": pog_errs}
    report["errors"].extend(pog_errs)

    # Stratified sample by account among ref rows with base ttl
    print(f"=== stratified sample (n≈{args.sample}) ===")
    by_acc = defaultdict(list)
    for ref in ref_rows:
        if ref.get("ttl") in (None, ""):
            continue
        by_acc[str(ref.get("acc") or "")].append(ref)
    rng = random.Random(args.seed)
    picked = []
    accounts = sorted(by_acc.keys(), key=lambda a: -len(by_acc[a]))
    # round-robin until sample size
    pointers = {a: 0 for a in accounts}
    for a in accounts:
        rng.shuffle(by_acc[a])
    while len(picked) < args.sample and any(pointers[a] < len(by_acc[a]) for a in accounts):
        for a in accounts:
            if pointers[a] < len(by_acc[a]):
                picked.append(by_acc[a][pointers[a]])
                pointers[a] += 1
                if len(picked) >= args.sample:
                    break
    sample_errs = []
    sample_ok = 0
    sample_miss = 0
    for ref in picked:
        key = biz_key(ref)
        g_cands = [
            g
            for g in gen_rows
            if str(g.get("acc") or "").strip() == key[0]
            and str(g.get("inv") or "").strip() == key[1]
            and D(g.get("ttl")) == key[2]
        ]
        if not g_cands:
            sample_miss += 1
            sample_errs.append(f"sample r{ref['row']} {key}: MISSING")
            continue
        errs = compare_pair(ref, g_cands[0], f"s_r{ref['row']}")
        if errs:
            sample_errs.extend(errs)
        else:
            sample_ok += 1
    print(f"  picked={len(picked)} ok={sample_ok} miss={sample_miss} diff={len(picked)-sample_ok-sample_miss}")
    if sample_errs[:8]:
        for e in sample_errs[:8]:
            print(" ", e)
        if len(sample_errs) > 8:
            print(f"  ... +{len(sample_errs) - 8} more")
    report["checks"]["stratified_sample"] = {
        "ok": not sample_errs,
        "picked": len(picked),
        "matched_ok": sample_ok,
        "errors": sample_errs,
    }
    report["errors"].extend(sample_errs)

    report["pass"] = (
        sum_ok and not ctrl_errs and not pog_errs and not sample_errs
    )
    if args.json:
        Path(args.json).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print("Wrote", args.json)

    print("RESULT:", "PASS" if report["pass"] else "FAIL")
    return 0 if report["pass"] else 1


if __name__ == "__main__":
    sys.exit(main())
