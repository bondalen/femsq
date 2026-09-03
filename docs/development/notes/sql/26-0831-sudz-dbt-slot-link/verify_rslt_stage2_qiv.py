#!/usr/bin/env python3
"""
Stage2 QIV gate: FEMSQ Rslt (base–QI–QII–QIV) vs Access ags_Yr_DbtChangesRslt_*.xlsx.

QIII в эталоне и в GEN игнорируется (колонки не сравниваются).

Usage:
  python3 verify_rslt_stage2_qiv.py --ref PATH.xlsx --gen PATH.xlsx [--json OUT.json]
  # exit 0 = PASS, 1 = FAIL

REF (Access): фиксированные 1-based колонки с ciaKey в базовом блоке.
GEN (FEMSQ): колонки срезов определяются по row3 (*_Ttl / *_Overd / *_погашено).
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import Counter, defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("NEED openpyxl (e.g. /tmp/rslt-venv/bin/python)", file=sys.stderr)
    sys.exit(2)

# Access Rslt (1-based); QIII/QIV — только QIV для stage2
REF = {
    "acc": 2,
    "inv": 3,
    "cn": 5,
    "base_ttl": 11,
    "base_o": 12,
    "q1_ttl": 25,
    "q1_o": 26,
    "q1_pog": 32,
    "q2_ttl": 40,
    "q2_o": 41,
    "q2_pog": 47,
    "q4_ttl": 71,
    "q4_o": 72,
    "q4_pog": 78,
}

SUM_KEYS = (
    "base_ttl",
    "base_o",
    "q1_ttl",
    "q1_o",
    "q1_pog",
    "q2_ttl",
    "q2_o",
    "q2_pog",
    "q4_ttl",
    "q4_o",
    "q4_pog",
)

def D(v):
    if v is None or v == "":
        return None
    try:
        return Decimal(str(v)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except Exception:
        return None


def m0(v):
    return D(v) if D(v) is not None else Decimal("0")


def load_rows(path: str, cols: dict, min_row: int = 4) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    max_c = max(cols.values())
    out = []
    for r, row in enumerate(
        ws.iter_rows(min_row=min_row, max_col=max_c, values_only=True), start=min_row
    ):
        if not row:
            continue
        rec = {"row": r}
        for name, idx in cols.items():
            rec[name] = row[idx - 1] if len(row) >= idx else None
        out.append(rec)
    wb.close()
    return out


def detect_gen_cols(path: str) -> dict:
    """Строит карту колонок GEN по технической строке (row3).

    FEMSQ export: срезы по uplDate (QIV часто ``2026-01-15``, Access — ``2026-01-30``).
    Берём первые три блока *_Ttl как base/QI/QII и **последний** как QIV (QIII пропускаем).
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    row3 = list(ws.iter_rows(min_row=3, max_row=3, values_only=True))[0]
    wb.close()

    ttl_cols: list[tuple[int, str]] = []
    for i, val in enumerate(row3, start=1):
        if val and str(val).endswith("_Ttl"):
            ttl_cols.append((i, str(val).replace("_Ttl", "")))

    if len(ttl_cols) < 4:
        raise ValueError(f"GEN: expected >=4 slice Ttl columns in {path}, got {len(ttl_cols)}")

    by_suffix: dict[str, int] = {}
    for i, val in enumerate(row3, start=1):
        if val is None or val == "":
            continue
        tech = str(val)
        for suffix in ("_Ttl", "_Overd", "_погашено"):
            if tech.endswith(suffix):
                by_suffix[tech] = i

    def pick(prefix: str, suffix: str) -> int:
        col = by_suffix.get(prefix + suffix)
        if col is None:
            raise ValueError(f"GEN column not found: {prefix}{suffix} in {path}")
        return col

    cols: dict[str, int] = {"acc": 2, "inv": 3}
    for i, val in enumerate(row3, start=1):
        if val and str(val).endswith("_cnNumEnum"):
            cols["cn"] = i
            break
    if "cn" not in cols:
        cols["cn"] = 5

    base_p, q1_p, q2_p = (p for _, p in ttl_cols[:3])
    q4_p = ttl_cols[-1][1]

    cols["base_ttl"] = pick(base_p, "_Ttl")
    cols["base_o"] = pick(base_p, "_Overd")
    cols["q1_ttl"] = pick(q1_p, "_Ttl")
    cols["q1_o"] = pick(q1_p, "_Overd")
    cols["q1_pog"] = pick(q1_p, "_погашено")
    cols["q2_ttl"] = pick(q2_p, "_Ttl")
    cols["q2_o"] = pick(q2_p, "_Overd")
    cols["q2_pog"] = pick(q2_p, "_погашено")
    cols["q4_ttl"] = pick(q4_p, "_Ttl")
    cols["q4_o"] = pick(q4_p, "_Overd")
    cols["q4_pog"] = pick(q4_p, "_погашено")
    cols["q4_slice_prefix"] = q4_p  # для отчёта

    required = set(SUM_KEYS) | {"acc", "inv", "cn"}
    missing = required - set(cols)
    if missing:
        raise ValueError(f"GEN columns incomplete in {path}: missing {sorted(missing)}")
    return cols


def bkey(r: dict) -> tuple:
    return (
        str(r.get("acc") or "").strip(),
        str(r.get("inv") or "").strip(),
        str(r.get("cn") or "").strip(),
        D(r.get("base_ttl")),
    )


def q4_key(r: dict) -> tuple:
    return (bkey(r), D(r.get("q4_ttl")), D(r.get("q4_o")), D(r.get("q4_pog")))


def sum_all(rows: list[dict]) -> dict[str, Decimal]:
    totals = {k: Decimal("0") for k in SUM_KEYS}
    for rec in rows:
        for k in SUM_KEYS:
            totals[k] += m0(rec.get(k))
    return totals


def compare_sums(ref_rows: list[dict], gen_rows: list[dict]) -> tuple[bool, dict, list[str]]:
    rt = sum_all(ref_rows)
    gt = sum_all(gen_rows)
    detail = {}
    errors = []
    ok_all = True
    for k in SUM_KEYS:
        ok = rt[k] == gt[k]
        ok_all = ok_all and ok
        detail[k] = {
            "ref": str(rt[k]),
            "gen": str(gt[k]),
            "delta": str(gt[k] - rt[k]),
            "ok": ok,
        }
        if not ok:
            errors.append(f"SUM {k}: ref={rt[k]} gen={gt[k]} delta={gt[k]-rt[k]}")
    return ok_all, detail, errors


def compare_q4_multiset(ref_rows: list[dict], gen_rows: list[dict]) -> tuple[bool, dict, list[str]]:
    rc = Counter(q4_key(r) for r in ref_rows)
    gc = Counter(q4_key(r) for r in gen_rows)
    extras = [(k, v - rc[k]) for k, v in gc.items() if v > rc[k]]
    missing = [(k, v - gc[k]) for k, v in rc.items() if v > gc[k]]
    extra_ttl = Decimal("0")
    for (bk, ttl, _o, _p), cnt in extras:
        extra_ttl += m0(ttl) * cnt
    info = {
        "extras": len(extras),
        "missing": len(missing),
        "extra_ttl_sum": str(extra_ttl),
    }
    errors = []
    if extras or missing:
        errors.append(f"QIV multiset extras={len(extras)} missing={len(missing)}")
    return not (extras or missing), info, errors


def top_q4_mismatches(ref_rows: list[dict], gen_rows: list[dict], limit: int = 15) -> list:
    idx_r = defaultdict(list)
    for r in ref_rows:
        idx_r[bkey(r)].append(r)
    idx_g = defaultdict(list)
    for g in gen_rows:
        idx_g[bkey(g)].append(g)

    mism = []
    for k, rr in idx_r.items():
        ref = rr[0]
        if k not in idx_g:
            mism.append(
                {
                    "kind": "missing_gen",
                    "key": list(k),
                    "ref_q4_ttl": str(D(ref.get("q4_ttl"))),
                    "ref_q4_pog": str(D(ref.get("q4_pog"))),
                }
            )
            continue
        gen = idx_g[k][0]
        if D(ref.get("q4_ttl")) != D(gen.get("q4_ttl")) or D(ref.get("q4_pog")) != D(
            gen.get("q4_pog")
        ):
            mism.append(
                {
                    "kind": "diff",
                    "key": list(k),
                    "ref_q4_ttl": str(D(ref.get("q4_ttl"))),
                    "gen_q4_ttl": str(D(gen.get("q4_ttl"))),
                    "ref_q4_pog": str(D(ref.get("q4_pog"))),
                    "gen_q4_pog": str(D(gen.get("q4_pog"))),
                }
            )
    mism.sort(
        key=lambda x: abs(
            m0(x.get("ref_q4_ttl")) - m0(x.get("gen_q4_ttl", x.get("ref_q4_ttl")))
        ),
        reverse=True,
    )
    return mism[:limit]


def _json_default(o):
    if isinstance(o, Decimal):
        return str(o)
    raise TypeError(f"Object of type {type(o).__name__} is not JSON serializable")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--ref", required=True, help="Access Rslt xlsx")
    ap.add_argument("--gen", required=True, help="FEMSQ Rslt xlsx (asOfUpl=901)")
    ap.add_argument("--json", default="", help="write report JSON")
    args = ap.parse_args()

    gen_cols = detect_gen_cols(args.gen)
    q4_prefix = gen_cols.pop("q4_slice_prefix", None)
    ref_rows = load_rows(args.ref, REF)
    gen_rows = load_rows(args.gen, gen_cols)

    report = {
        "pass": True,
        "scope": "base-QI-QII-QIV (QIII ignored)",
        "rows": {"ref": len(ref_rows), "gen": len(gen_rows)},
        "gen_cols": {**gen_cols, "q4_slice_prefix": q4_prefix},
        "sums": {},
        "errors": [],
    }

    sum_ok, sum_detail, sum_errs = compare_sums(ref_rows, gen_rows)
    report["sums"] = sum_detail
    report["errors"].extend(sum_errs)

    ms_ok, ms_info, ms_errs = compare_q4_multiset(ref_rows, gen_rows)
    report["q4_multiset"] = ms_info
    report["errors"].extend(ms_errs)

    report["top_q4_mismatches"] = top_q4_mismatches(ref_rows, gen_rows)
    report["pass"] = sum_ok and ms_ok

    print("=== row1 SUBTOTAL base–QI–QII–QIV ===")
    for k in SUM_KEYS:
        d = sum_detail[k]
        mark = "OK" if d["ok"] else "DIFF"
        print(f"  {k}: ref={d['ref']} gen={d['gen']} delta={d['delta']} {mark}")

    print("=== QIV multiset ===")
    print(
        f"  extras={ms_info['extras']} missing={ms_info['missing']} "
        f"extra_ttl_sum={ms_info['extra_ttl_sum']}"
    )
    print(f"  top mismatches: {len(report['top_q4_mismatches'])}")

    if args.json:
        Path(args.json).write_text(
            json.dumps(
                report,
                ensure_ascii=False,
                indent=2,
                default=_json_default,
            ),
            encoding="utf-8",
        )
        print("Wrote", args.json)

    print("RESULT:", "PASS" if report["pass"] else "FAIL")
    return 0 if report["pass"] else 1


if __name__ == "__main__":
    sys.exit(main())
